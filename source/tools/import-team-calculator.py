#!/usr/bin/env python3
"""
IMPORT THE TEAM'S CALCULATOR WORKBOOK INTO THE SITE (D-144, 25 Sep 2026).

Owner's ruling 1 (24 Sep 2026, late night): wire the agronomy team's calculator workbook into the site
for the candidate crops, the way wheat and potato are wired, and skip any crop that fails a check.

Re-run this whenever the team sends a new workbook:

    python3 tools/import-team-calculator.py ["/path/to/Fertilizer Nutrient Calculator ....xlsx"]

then `npx tsc -p tsconfig.enginetest.json && node run-enginetest.cjs __teamcalccheck`, then the build.

WHAT IT DOES
  1. Reads every sheet with openpyxl (formulas, not cached values: the workbook ships with none). A
     quantity typed as a formula ("=3/4", "=(1/4)+(1/2)") is evaluated by a small arithmetic parser
     that accepts digits, + - * / ( ) and nothing else.
  2. Maps each workbook product name to a site product (NAME_MAP). A name with no site product
     ("Zinc Coated Urea", "FeSO4", the 12-44 "Ammonium Phosphate") is NOT invented: the row is
     reported as not counted and the crop is skipped.
  3. Takes analysis and pack from the SITE, never from the sheet: analysis from MASTER_PRODUCTS in
     src/data/pricing.ts (the owner's confirmed figures, e.g. Fusion Potash K2O 50, N 2, B 1, S 1.5,
     Mg 1), pack from the crop's own published plan in src/data/catalogue.ts. Every place the sheet
     disagrees is logged. The sheet's price column is never read: prices come from FARMER_PRICE in
     pricing.ts, and only after the farmer presses the cost button, as for wheat and potato.
  4. Applies the owner's rulings that the workbook has not caught up with (RULING_OVERRIDES).
  5. Checks every candidate crop and SKIPS it, with the reason, if any check fails:
       - the sheet maps to a site crop that has a published plan
       - every row with a quantity maps to a site product
       - every quantity parses to a positive, finite number
       - per-acre nutrient totals are inside sane bounds (SANE)
       - no per-tree or age-band mixing (mango and date palm are excluded outright; the quantity
         header must say "per acre"; the site plan must carry no per-plant rows and no age bands)
       - no per-hectare units anywhere on the sheet
     The sheet's own "Recommended Dose" band is extracted and printed in the report for comparison.
     The site never prints it and never claims anything about it.
  6. Writes src/data/teamCalculators.ts (generated, do not edit by hand) and
     tools/team-calculator-import-report.txt.
"""
import datetime
import json
import math
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = '/mnt/user-data/uploads/VAN Web APP/Output 24 Sep 2026/Fertilizer Nutrient Calculator For Economic Nutrition Planning.xlsx'
OUT_TS = os.path.join(ROOT, 'src', 'data', 'teamCalculators.ts')
OUT_REPORT = os.path.join(ROOT, 'tools', 'team-calculator-import-report.txt')

# The candidate crops named in the owner's ruling 1 (from /root/van/audit/team-tools-review.md, wave 1).
# sheet name -> site crop slug. Order here is the order the crops are listed on the site.
CANDIDATES = [
    ('Sugarcane (Ratoon)', 'sugarcane-ratoon'),
    ('Cotton', 'cotton'),
    ('Maize', 'maize'),
    ('Rice (Basmati)', 'rice-basmati'),
    ('Canola', 'canola'),
    ('Sesame', 'sesame'),
    ('Chickpea', 'chickpea'),
    ('Lentil', 'lentil'),
    ('Chili', 'chili'),
    ('Garlic', 'garlic'),
    ('Citrus (5 years)', 'citrus'),
]
# Never wired from this workbook: per-tree or age-band programmes (the owner's rule: no per-tree and
# per-acre mixing). Listed so the report says so rather than leaving them silently out.
EXCLUDED = {'mango': 'per-tree programme (the sheet mixes per-plant and per-acre quantities)',
            'date-palm': 'age-band programme (the sheet adds the 5 tree-age bands into 1 figure)'}
ALREADY_LIVE = {'Wheat': 'wheat', 'Potato': 'potato'}

# workbook name (trimmed) -> MASTER_PRODUCTS name in pricing.ts, or a commodity rule.
# 'plan:CAN' means: use the crop plan's own name and printed analysis for its calcium ammonium nitrate
# row (the site does not attribute a third-party brand's analysis to a generic material).
NAME_MAP = {
    'Vital Urea': 'Vital Urea',
    'Urea': 'Urea',
    'Ammonium Sulphate': 'Ammonium Sulphate',
    'V-Ammonium Phosphate': 'V. Ammonium Phosphate (V-Phosphate)',
    'Sarsabz CAN': 'plan:CAN',
    'Green Phosphate': 'Green Phosphate',
    'Fusion Phosphate': 'Fusion Phosphate',
    'Vital Potash': 'Vital Potash',
    'Fusion Potash': 'Fusion Potash',
    'SOP': 'SOP',
    'V-Germinator Pro': 'V-Germinator Pro',
    'VL-Potash': 'VL-Potash',
    'V-Potash Plus': 'V-Potash Plus',
    'Crop Force (15:15:15)': 'Crop Force (15:15:15)',
    'Crop Force (12:12:18)': 'Crop Force (12:12:18)',
    'VL-NPK (8:8:6)': 'VL-NPK (8:8:6)',
    'VL-Micromix': 'VL-Micromix',
    'VL-Boron': 'VL-Boron',
    'V-Zinc 10%': 'V-Zinc 10%',
    'V-Transform': 'V-Transform',
    'Cala-Mag V': 'Cala-Mag V',
    'V-Mag Essential': 'V-Mag Essential',
    'Green Sulfur': 'Green Sulfur',
    # 2 grades on the site (owner, 8 Sep 2026): 20 kg Mn grade and 25 kg OM grade. Resolved by the plan's pack.
    'V-Compost': 'slug:v-compost',
    'Humi Grow': 'Humi Grow',
    'Humi Grow Plus': 'Humi Grow Plus',
    'Tornado': 'Tornado',
}
# Names in the workbook with no site product. Never invented.
UNMAPPED = {
    'Zinc Coated Urea': 'no zinc-coated grade exists at VAN (D-100) and the sheet gives it no analysis or pack',
    'FeSO4': 'not a VAN product and the sheet gives it no analysis',
    'Ammonium Phosphate': 'the 12-44 grade is out of date; the owner ruled V-Ammonium Phosphate is 10-44-0 (8 Sep 2026)',
    'Sarsabz NP': 'a third-party product in no VAN plan',
    'DAP': 'a commodity in no VAN plan', 'SSP': 'a commodity in no VAN plan',
    'TSP': 'a commodity in no VAN plan', 'MOP': 'a commodity in no VAN plan',
}

# Owner's rulings the workbook has not caught up with. (crop slug, workbook name) -> what to use.
RULING_OVERRIDES = {
    ('garlic', 'Crop Force (15:15:15)'): {
        'qtyPerAcre': 0.25, 'packKg': 20,
        'plan_rate': '¼ bag', 'plan_pack': 'Bag - 20 kg',
        'why': 'Owner, 10 Sep 2026: "Garlic 5 kg, is 1/4 bag of same 20 kg bag." The workbook still has 1 x 1 kg. Site figure used (0.25 x 20 kg = 5 kg).',
    },
}

# Sane per-acre bounds, kg. Wider than any published Pakistani recommendation for these crops; a total
# outside them means a transcription or unit error, not agronomy.
SANE = {'N': (1.0, 200.0), 'P': (0.0, 150.0), 'K': (0.0, 150.0)}
MAX_ROW_KG = 250.0          # no single product row above 250 kg or L an acre
PER_HA_RE = re.compile(r'hectare|\bper\s*ha\b|kg\s*/\s*ha\b|t\s*/\s*ha\b|/\s*ha\b', re.I)


def safe_eval(expr):
    """Evaluate '=3/4', '=(1/4)+(1/2)', 1.5. Only digits, . + - * / ( ) and spaces are accepted."""
    if isinstance(expr, (int, float)):
        return float(expr)
    s = str(expr).strip()
    if s.startswith('='):
        s = s[1:]
    if not re.fullmatch(r'[0-9.+\-*/() ]+', s):
        raise ValueError(f'not plain arithmetic: {expr!r}')
    tokens = re.findall(r'\d+\.?\d*|\.\d+|[+\-*/()]', s)
    pos = 0

    def peek():
        return tokens[pos] if pos < len(tokens) else None

    def take():
        nonlocal pos
        t = tokens[pos]; pos += 1; return t

    def factor():
        t = peek()
        if t == '(':
            take(); v = expr_(); assert take() == ')'; return v
        if t == '-':
            take(); return -factor()
        return float(take())

    def term():
        v = factor()
        while peek() in ('*', '/'):
            op = take(); r = factor(); v = v * r if op == '*' else v / r
        return v

    def expr_():
        v = term()
        while peek() in ('+', '-'):
            op = take(); r = term(); v = v + r if op == '+' else v - r
        return v

    v = expr_()
    if pos != len(tokens):
        raise ValueError(f'unparsed tail in {expr!r}')
    return v


def load_master():
    src = open(os.path.join(ROOT, 'src', 'data', 'pricing.ts'), encoding='utf-8').read()
    block = src[src.index('export const MASTER_PRODUCTS'):]
    block = block[:block.index('\n]\n')]
    out = []
    for m in re.finditer(r'\{ product: "([^"]+)", slug: (null|"([^"]+)"), analysisPct: \{([^}]*)\}, packKg: ([\d.]+)', block):
        analysis = {}
        for k, v in re.findall(r'(\w+): ([\d.]+)', m.group(4)):
            analysis[k] = float(v)
        out.append({'product': m.group(1), 'slug': m.group(3), 'analysisPct': analysis, 'packKg': float(m.group(5))})
    return out


def load_plans():
    """Site crop plans from catalogue.ts: {slug: {'rows': [...], 'mode': 'age'|None}}."""
    src = open(os.path.join(ROOT, 'src', 'data', 'catalogue.ts'), encoding='utf-8').read()
    consts = {}
    for m in re.finditer(r'export const (\w+_PLAN): PlanRow\[\] = \[(.*?)\n\]', src, re.S):
        body = re.sub(r'/\*.*?\*/', '', m.group(2), flags=re.S)
        body = '\n'.join(l for l in body.split('\n') if not l.strip().startswith('//'))
        body = re.sub(r',\s*([\]}])', r'\1', '[' + body.strip().rstrip(',') + ']')
        consts[m.group(1)] = json.loads(body)
    plans = {}
    block = src[src.index('export const CROP_PLANS'):]
    for m in re.finditer(r'"([\w-]+)": \{ stages: \w+, plan: (\w+)(?:, mode: "(\w+)")? \}', block):
        plans[m.group(1)] = {'rows': consts[m.group(2)], 'mode': m.group(3)}
    return plans


def fmt_an(a):
    return ', '.join(f'{k} {v:g}' for k, v in a.items()) or 'none'


def pack_of(pack_str):
    m = re.search(r'([\d.]+)\s*(kg|L|g)\b', pack_str)
    if not m:
        return None, None
    v = float(m.group(1))
    return (v / 1000 if m.group(2) == 'g' else v), ('L' if m.group(2) == 'L' else 'kg')


def plan_analysis(text):
    """Parse a plan row's printed analysis for a commodity ('N 26%', 'N 21% · S 24%')."""
    out = {}
    for k, v in re.findall(r'\b(N|S|K|Ca|Mg)\s+([\d.]+)%', text):
        out[k] = float(v)
    return out


def sheet_columns(ws):
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        if v is not None:
            hdr[str(v).strip()] = c
    qty_col = next((c for k, c in hdr.items() if k.startswith('QTY')), None)
    qty_hdr = next((k for k in hdr if k.startswith('QTY')), '')
    return hdr, qty_col, qty_hdr


def read_sheet(ws):
    hdr, qty_col, qty_hdr = sheet_columns(ws)
    pct_cols = {k.replace('(%)', '').strip(): c for k, c in hdr.items() if k.endswith('(%)')}
    pack_col = next(c for c in range(1, ws.max_column + 1) if str(ws.cell(1, c).value or '').strip().lower().startswith('net weight'))
    rows, problems = [], []
    last_product_row = 2
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name is None or str(name).strip() in ('Gross Total',):
            continue
        last_product_row = r
        raw = ws.cell(r, qty_col).value if qty_col else None
        if raw is None or str(raw).strip() == '':
            continue
        try:
            qty = safe_eval(raw)
        except Exception as e:  # noqa: BLE001
            problems.append(f'row {r} {str(name).strip()}: quantity {raw!r} does not parse ({e})')
            continue
        analysis = {}
        for k, c in pct_cols.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and v:
                analysis[k] = float(v)
        pack = ws.cell(r, pack_col).value
        rows.append({'row': r, 'name': str(name).strip(), 'raw': raw, 'qty': qty, 'sheetAnalysis': analysis,
                     'sheetPack': float(pack) if isinstance(pack, (int, float)) else None})
    # Recommended dose rows: the row whose quantity cell reads "Recommended Dose", then labelled rows.
    bands = []
    tot = {k: hdr.get(k) for k in ('Total N', 'Total P', 'Total K')}
    for r in range(3, ws.max_row + 1):
        if qty_col and str(ws.cell(r, qty_col).value or '').strip().lower() == 'recommended dose':
            for rr in range(r, ws.max_row + 1):
                vals = [ws.cell(rr, tot[k]).value if tot[k] else None for k in ('Total N', 'Total P', 'Total K')]
                if all(isinstance(v, (int, float)) for v in vals):
                    label = ws.cell(rr, qty_col).value
                    label = '' if label is None or str(label).strip().lower() == 'recommended dose' else str(label).strip()
                    bands.append({'label': label or 'Recommended', 'N': float(vals[0]), 'P': float(vals[1]), 'K': float(vals[2])})
            break
    # Per-hectare scan: every text cell, every formula (a /2.471 conversion is a per-hectare figure).
    per_ha = []
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and (PER_HA_RE.search(v) or re.search(r'2\.47\d*', v)):
                per_ha.append(f'{c.coordinate}: {v[:60]}')
    return {'rows': rows, 'problems': problems, 'bands': bands, 'qtyHeader': qty_hdr, 'perHa': per_ha}


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx)  # formulas, not cached values
    master = load_master()
    mbyname = {m['product']: m for m in master}
    plans = load_plans()

    report = []
    live, skipped = {}, []
    P = report.append
    P('TEAM CALCULATOR IMPORT REPORT')
    P(f'Workbook: {xlsx}')
    P(f'Sheets: {len(wb.sheetnames)}. Run: {datetime.date.today().isoformat()}. Script: tools/import-team-calculator.py')
    P('Analysis and pack: the site\'s (MASTER_PRODUCTS, crop plan). Prices: never read from the sheet.')
    P('')

    sheet_to_slug = dict(CANDIDATES)
    for name in wb.sheetnames:
        if name not in sheet_to_slug and name not in ALREADY_LIVE:
            P(f'-- {name}: not a candidate in ruling 1, not wired.')
    for s, why in EXCLUDED.items():
        P(f'-- {s}: excluded outright, {why}.')
    P('')

    for sheet, slug in CANDIDATES:
        log, fails = [], []
        P(f'=== {sheet} -> {slug} ===')
        if sheet not in wb.sheetnames:
            fails.append('sheet not found in the workbook')
            data = {'rows': [], 'problems': [], 'bands': [], 'qtyHeader': '', 'perHa': []}
        else:
            data = read_sheet(wb[sheet])
        plan = plans.get(slug)
        if not plan:
            fails.append('no published site plan for this crop')
        if slug in EXCLUDED:
            fails.append(EXCLUDED[slug])
        if plan and plan['mode'] == 'age':
            fails.append('site plan is by tree age band')
        if plan and any(r.get('basis') == 'plant' for r in plan['rows']):
            fails.append('site plan carries per-plant rows')
        qh = data['qtyHeader'].lower()
        if 'per acre' not in qh or 'plant' in qh or 'tree' in qh:
            fails.append(f'quantity column is not per acre: "{data["qtyHeader"]}"')
        if data['perHa']:
            fails.append('per-hectare units on the sheet: ' + '; '.join(data['perHa'][:3]))
        fails.extend(data['problems'])

        out_rows = []
        sheet_tot = {'N': 0.0, 'P': 0.0, 'K': 0.0}
        site_tot = {'N': 0.0, 'P': 0.0, 'K': 0.0}
        plan_rows = plan['rows'] if plan else []
        for w in data['rows']:
            nm = w['name']
            for k in ('N', 'P', 'K'):
                sheet_tot[k] += w['qty'] * (w['sheetPack'] or 0) * w['sheetAnalysis'].get(k, 0) / 100
            if w['qty'] <= 0 or not math.isfinite(w['qty']):
                fails.append(f'row {w["row"]} {nm}: quantity {w["raw"]!r} is not a positive number')
                continue
            if nm in UNMAPPED or nm not in NAME_MAP:
                why = UNMAPPED.get(nm, 'no site product carries this name')
                fails.append(f'row {w["row"]} "{nm}" maps to no site product ({why}); not counted')
                continue
            target = NAME_MAP[nm]
            note = []
            if target == 'plan:CAN':
                can = [r for r in plan_rows if r['commodity'] and re.search(r'\bCAN\b|calcium ammonium nitrate', r['product'], re.I)]
                if not can:
                    fails.append(f'row {w["row"]} {nm}: the site plan has no CAN row to map it to')
                    continue
                product, slug_, analysis = can[0]['product'], None, plan_analysis(can[0]['analysis'])
                pack, _ = pack_of(can[0]['pack'])
                note.append(f'"{nm}" shown as the plan\'s "{product}" with the plan\'s printed analysis ({fmt_an(analysis)}); the sheet has {fmt_an(w["sheetAnalysis"])}.')
            else:
                m = next((x for x in master if x['slug'] == target[5:]), None) if target.startswith('slug:') else mbyname.get(target)
                if not m:
                    fails.append(f'row {w["row"]} {nm}: MASTER_PRODUCTS has no "{target}"')
                    continue
                slug_ = m['slug']
                same_slug = [r for r in plan_rows if (r['slug'] == slug_ if slug_ else (not r['slug'] and r['product'].lower().replace('sulfate', 'sulphate') == target.lower()))]
                plan_packs = sorted({pack_of(r['pack'])[0] for r in same_slug if pack_of(r['pack'])[0]})
                product = same_slug[0]['product'] if same_slug else m['product']
                if w['sheetPack'] in plan_packs:
                    pack = w['sheetPack']
                elif len(plan_packs) == 1:
                    pack = plan_packs[0]
                    note.append(f'pack: sheet {w["sheetPack"]:g}, site plan {pack:g}; site used.')
                elif not plan_packs and m['packKg'] == w['sheetPack']:
                    pack = m['packKg']
                elif not plan_packs:
                    pack = m['packKg']
                    note.append(f'pack: sheet {w["sheetPack"]:g}, site catalogue {pack:g}; site used.')
                else:
                    fails.append(f'row {w["row"]} {nm}: sheet pack {w["sheetPack"]} is none of the site plan packs {plan_packs}')
                    continue
                # Analysis: the MASTER record for this slug at this pack if there is one (V-Compost has 2
                # grades by pack), otherwise the mapped record.
                by_pack = [x for x in master if slug_ and x['slug'] == slug_ and abs(x['packKg'] - pack) < 1e-9]
                rec = by_pack[0] if len(by_pack) == 1 else m
                analysis = dict(rec['analysisPct'])
                if rec is not m or target.startswith('slug:'):
                    note.append(f'analysis from the site\'s {rec["product"]} record.')
            ov = RULING_OVERRIDES.get((slug, nm))
            qty = w['qty']
            if ov:
                match = [r for r in plan_rows if r['slug'] == slug_ and r['rate'] == ov['plan_rate'] and r['pack'] == ov['plan_pack']]
                if not match:
                    fails.append(f'row {w["row"]} {nm}: ruling override no longer matches the site plan row; re-check the ruling')
                    continue
                note.append(f'sheet {qty:g} x {w["sheetPack"]:g}; ' + ov['why'])
                qty, pack = ov['qtyPerAcre'], ov['packKg']
            sa = {k: v for k, v in w['sheetAnalysis'].items() if k not in ('HA', 'OM', 'AA') or k in analysis}
            if target != 'plan:CAN':
                diff = sorted(set(sa) | set(analysis))
                d = [f'{k} sheet {sa.get(k, 0):g} site {analysis.get(k, 0):g}' for k in diff if abs(sa.get(k, 0) - analysis.get(k, 0)) > 1e-9]
                if d:
                    note.append('analysis differs, site used: ' + ', '.join(d) + '.')
            kg = qty * pack
            if kg > MAX_ROW_KG:
                fails.append(f'row {w["row"]} {nm}: {kg:g} kg an acre in one row is outside sane bounds')
            for k in ('N', 'P', 'K'):
                site_tot[k] += kg * analysis.get(k, 0) / 100
            if note:
                log.append(f'  row {w["row"]} {nm}: ' + ' '.join(note))
            out_rows.append({'product': product, 'slug': slug_, 'analysisPct': analysis, 'packKg': pack,
                             'qtyPerAcre': round(qty, 6), 'unitPricePkr': 0, 'sheetRow': w['row'],
                             'sheetName': nm, 'sheetQty': str(w['raw']), 'note': ' '.join(note)})
        for k, (lo, hi) in SANE.items():
            if not (lo <= site_tot[k] <= hi):
                fails.append(f'{k} total {site_tot[k]:.1f} kg an acre is outside the sane bounds {lo:g} to {hi:g}')
        if not out_rows:
            fails.append('no rows')

        P(f'  rows with a quantity: {len(data["rows"])}   mapped: {len(out_rows)}')
        P(f'  quantity header: "{data["qtyHeader"]}"')
        P('  sheet totals (sheet analysis x sheet pack), kg an acre: N {N:.2f} · P2O5 {P:.2f} · K2O {K:.2f}'.format(**sheet_tot))
        P('  site totals (site analysis x site pack), kg an acre:   N {N:.2f} · P2O5 {P:.2f} · K2O {K:.2f}'.format(**site_tot))
        for b in data['bands']:
            P(f'  recommended dose on the sheet, "{b["label"]}": N {b["N"]:g} · P2O5 {b["P"]:g} · K2O {b["K"]:g}  (report only; never printed on the site)')
        if not data['bands']:
            P('  recommended dose on the sheet: none')
        for x in log:
            P(x)
        if fails:
            P('  SKIPPED: ' + ' | '.join(fails))
            skipped.append({'sheet': sheet, 'slug': slug, 'reasons': fails})
        else:
            P('  WIRED: every check passed.')
            live[slug] = {'slug': slug, 'sheet': sheet, 'rows': out_rows, 'bands': data['bands'],
                          'sheetTotals': {k: round(v, 3) for k, v in sheet_tot.items()}}
        P('')

    P(f'WIRED {len(live)}: ' + ', '.join(live))
    P(f'SKIPPED {len(skipped)}: ' + ', '.join(f'{s["slug"]} ({s["reasons"][0]})' for s in skipped))

    ts = []
    ts.append('/**')
    ts.append(' * GENERATED by tools/import-team-calculator.py. Do not edit by hand: re-run the script when the team')
    ts.append(' * sends a new workbook. D-144 (owner\'s ruling 1, 24 Sep 2026).')
    ts.append(f' * Source: {os.path.basename(xlsx)}')
    ts.append(' * Analysis and pack are the site\'s own (MASTER_PRODUCTS, the crop plan). The sheet\'s prices are never')
    ts.append(' * read: unitPricePkr is 0 here and the cost panel prices each line from FARMER_PRICE after the click.')
    ts.append(' * Each row keeps the sheet row, name and quantity it came from, and a note wherever the site\'s figure')
    ts.append(' * replaced the sheet\'s. The recommended-dose bands are for the engine report only; no page prints them.')
    ts.append(' */')
    ts.append("import type { CostRow } from './pricing'")
    ts.append('')
    ts.append('export type TeamCalcRow = CostRow & { sheetRow: number; sheetName: string; sheetQty: string; note: string }')
    ts.append('export type TeamCalcBand = { label: string; N: number; P: number; K: number }')
    ts.append('export type TeamCalculator = { slug: string; sheet: string; rows: TeamCalcRow[]; bands: TeamCalcBand[]; sheetTotals: { N: number; P: number; K: number } }')
    ts.append('')
    ts.append(f'export const TEAM_CALC_SOURCE = {json.dumps({"file": os.path.basename(xlsx), "sheets": len(wb.sheetnames), "imported": datetime.date.today().isoformat()}, ensure_ascii=False)}')
    ts.append('')
    ts.append('/** Crops that passed every check, in the order they are listed on the site. */')
    def num(v):
        return int(v) if isinstance(v, float) and v.is_integer() else v

    def compact(o):
        if isinstance(o, dict):
            return {k: compact(v) for k, v in o.items()}
        if isinstance(o, list):
            return [compact(v) for v in o]
        return num(o) if isinstance(o, float) else o

    ts.append('export const TEAM_CALCULATORS: Record<string, TeamCalculator> = {')
    for slug_, c in live.items():
        ts.append(f'  {json.dumps(slug_)}: {{')
        ts.append(f'    slug: {json.dumps(slug_)}, sheet: {json.dumps(c["sheet"], ensure_ascii=False)},')
        ts.append(f'    sheetTotals: {json.dumps(compact(c["sheetTotals"]))},')
        ts.append(f'    bands: {json.dumps(compact(c["bands"]), ensure_ascii=False)},')
        ts.append('    rows: [')
        for r in c['rows']:
            ts.append('      ' + json.dumps(compact(r), ensure_ascii=False) + ',')
        ts.append('    ],')
        ts.append('  },')
    ts.append('}')
    ts.append('')
    ts.append('/** Candidate crops that failed a check, with every reason. Not wired. */')
    ts.append('export const TEAM_CALC_SKIPPED: { sheet: string; slug: string; reasons: string[] }[] = ' + json.dumps(skipped, ensure_ascii=False, indent=1))
    ts.append('')
    open(OUT_TS, 'w', encoding='utf-8').write('\n'.join(ts))
    open(OUT_REPORT, 'w', encoding='utf-8').write('\n'.join(report) + '\n')
    print('\n'.join(report))
    print(f'\nwrote {os.path.relpath(OUT_TS, ROOT)} and {os.path.relpath(OUT_REPORT, ROOT)}')


if __name__ == '__main__':
    main()
